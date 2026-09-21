#!/usr/bin/env python3
"""Export both C. elegans connectomes (Cook et al. 2019) as species packages.

    .venv/bin/python tools/export_celegans.py            # fetch if missing, export both
    .venv/bin/python tools/export_celegans.py --no-fetch # raw files must be present

Raw files in data/raw/cook2019/ and data/raw/wang2024/ (gitignored), packages
out in web/species/c_elegans_herm/ and web/species/c_elegans_male/, and the
`export` block of both roster entries in data/roster.json rewritten with the
URL and sha256 of every raw file read.

WHAT COMES FROM WHERE

  neurons     Cook et al. 2019, SI 4 "Cell lists": every cell whose type is
              sensory, interneuron, motorneuron or neuron (CANL/R), in the
              sheet's own order. 302 hermaphrodite, 385 male.
  chemical    SI 5 "Connectome adjacency matrices, corrected July 2020",
              rows presynaptic. The weight is Cook's: the number of EM serial
              sections of connectivity (number AND size of synapses, gaps
              filled by extrapolation) -- not a count of synapses. It is stored
              as the i16 "count". Only neuron -> neuron edges are kept; edges
              onto muscles and other end organs are counted and left out.
  gap         SI 5 "gap jn symmetric", pairs i < j between neurons, weight in
              the same serial-section unit. Kept whatever their sign, because
              a gap junction has none.
  signs       Wang et al. 2024 (eLife 13:RP95402), Supplementary files 2
              (hermaphrodite) and 3 (male-specific neurons), column
              "Neurotransmitter(s)", through signs.transmitter_sign.
  sensory     SI 4's own modality notes -- see SENSORY below.
  motor       SI 4 cell type "motorneuron", dealt by package.deal_groups.
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import sys
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from open_animal_stage.signs import Sign, transmitter_sign          # noqa: E402
from open_animal_stage import package as pk                        # noqa: E402

RAW_COOK = os.path.join(ROOT, "data", "raw", "cook2019")
RAW_WANG = os.path.join(ROOT, "data", "raw", "wang2024")

WW = "https://wormwiring.org/si/"
COOK_FILES = {
    "cell_lists": "SI 4 Cell lists.xlsx",
    "adjacency": "SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",
}
WANG = "https://cdn.elifesciences.org/articles/95402/"
WANG_FILES = {
    "herm_nt": "elife-95402-supp2-v1.xlsx",
    "male_nt": "elife-95402-supp3-v1.xlsx",
    "dimorphic_nt": "elife-95402-supp4-v1.xlsx",
}

# Cook et al. 2019, main text: "The respective graphs have 4,887 chemical (or
# directed) edges and 1,447 gap junction (or undirected) edges in the
# hermaphrodite and 5,315 chemical and 1,755 gap junction edges in the male",
# over 460 / 579 nodes including muscles and end organs.
PAPER = {
    "herm": {"neurons": 302, "chem_edges": 4887, "gap_edges": 1447},
    "male": {"neurons": 385, "chem_edges": 5315, "gap_edges": 1755},
}
# The July 2020 correction of SI 5 post-dates the paper; edge totals may move a
# little. Neuron counts must match exactly.
EDGE_TOLERANCE = 0.01

LICENCE = (
    "Connectome: Cook et al. (2019) Nature 571:63-71, doi:10.1038/s41586-019-1352-7, "
    "Supplementary Information 4 and 5 as distributed by WormWiring.org "
    "(Emmons lab; page footer 'Copyright (c) 2020'). No open licence is stated on "
    "the files or the site; the paper is a Springer Nature article (NIH author "
    "manuscript PMC6889226). Treat the derived export as NOT cleared for "
    "redistribution until the authors are asked. Neurotransmitter atlas: Wang et al. "
    "(2024) eLife 13:RP95402, doi:10.7554/eLife.95402, CC-BY 4.0.")

# The four channels. Each is a set of cells the source itself annotates with
# that modality (SI 4, column "notes, sensilla, modality"); the check below
# fails if the annotation in the file does not say so.
SENSORY = {
    "attractive_chemosensory": {
        "signal": "reward",
        "neurons": ["ASEL", "ASER", "AWAL", "AWAR", "AWCL", "AWCR"],
        "annotation": ("sensory", "amphid"),
        "why": ("Amphid chemosensory neurons that mediate attraction: ASE to "
                "water-soluble attractants (salts), AWA and AWC to volatile "
                "attractants (Bargmann & Horvitz 1991; Bargmann et al. 1993). "
                "The closest thing the worm has to the fly's sugar neurons. SI 4 "
                "annotates them 'sensory, SN6, amphid'; SI 4 does not itself say "
                "'attractive' -- that part is the cited literature."),
    },
    "nociceptive": {
        "signal": "harm",
        "neurons": ["ASHL", "ASHR", "ADLL", "ADLR"],
        "annotation": ("sensory", "nociceptive"),
        "why": ("SI 4 annotates exactly these four 'amphid, nociceptive' (SN5): "
                "the avoidance neurons, the analogue of the fly's bitter channel."),
    },
    "o2_co2": {
        "signal": "reserve",
        "neurons": ["URXL", "URXR", "BAGL", "BAGR"],
        "annotation": ("sensory", "O2, CO2"),
        "why": ("SI 4 annotates URX and BAG 'O2, CO2, social signals, touch' "
                "(SN4). O2 (URX) and CO2 (BAG) sensing is the worm's homeostatic "
                "read of its surroundings' gas balance, standing in for the fly's "
                "water (thirst) channel, which the game drives from treasury. AQR "
                "and PQR are also O2 sensors in the literature but SI 4 annotates "
                "them 'touch', so they are not used: the source's label decides."),
    },
    "gentle_touch": {
        "signal": "threat",
        "neurons": ["ALML", "ALMR", "AVM", "PLML", "PLMR", "PVM"],
        "annotation": ("sensory", "mechanosensory"),
        "why": ("The six gentle-touch receptor neurons (Chalfie et al. 1985), "
                "annotated 'mechanosensory' (SN3) in SI 4: body mechanosensation, "
                "as the fly's Johnston's organ channel is."),
    },
}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def fetch(url, path, allow):
    if os.path.exists(path):
        return
    if not allow:
        raise SystemExit("missing %s (run without --no-fetch, or download %s)" % (path, url))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    print("  fetch", url)
    req = urllib.request.Request(url, headers={"User-Agent": "open-animal-stage-exporter"})
    with urllib.request.urlopen(req, timeout=120) as r, open(path + ".part", "wb") as f:
        f.write(r.read())
    os.replace(path + ".part", path)


def raw_files(allow_fetch):
    out = {}
    for key, name in COOK_FILES.items():
        url = WW + urllib.parse.quote(name)
        p = os.path.join(RAW_COOK, name)
        fetch(url, p, allow_fetch)
        out[key] = {"path": os.path.relpath(p, ROOT), "url": url, "sha256": sha256(p)}
    for key, name in WANG_FILES.items():
        url = WANG + name
        p = os.path.join(RAW_WANG, name)
        fetch(url, p, allow_fetch)
        out[key] = {"path": os.path.relpath(p, ROOT), "url": url, "sha256": sha256(p)}
    return out


def rows_of(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]


def clean(x):
    return str(x).strip() if x is not None else ""


# ------------------------------------------------------------------ neurons

NEURON_TYPES = {"sensory", "interneuron", "motorneuron", "neuron",
                "sensory neuron"}


def cell_list(path, sex):
    """[(name, type, notes)] for the neurons of one sex, in SI 4's order."""
    out = []
    for r in rows_of(path, "pharynx"):
        if clean(r[1]) in NEURON_TYPES:
            out.append((clean(r[0]), clean(r[1]), "pharynx"))
    for r in rows_of(path, "sex-shared")[1:]:
        if clean(r[1]) in NEURON_TYPES:
            out.append((clean(r[0]), clean(r[1]), clean(r[3]) + "|" + clean(r[2])))
    sheet = "hermaphrodite specific" if sex == "herm" else "male-specific"
    for r in rows_of(path, sheet):
        if clean(r[1]) in NEURON_TYPES:
            out.append((clean(r[0]), clean(r[1]), "sex-specific"))
    return out


def matrix(path, sheet):
    """{(row_name, col_name): value} for every numeric cell of an SI 5 sheet."""
    rows = rows_of(path, sheet)
    cols = {j: clean(c) for j, c in enumerate(rows[2]) if c is not None and j >= 3}
    out, odd = {}, []
    for r in rows[3:]:
        if len(r) < 3 or r[2] is None:
            continue
        a = clean(r[2])
        for j, b in cols.items():
            v = r[j] if j < len(r) else None
            if v is None or v == "":
                continue
            if not isinstance(v, (int, float)):
                odd.append((a, b, v))
                continue
            if v:
                out[(a, b)] = v
    if odd:
        raise SystemExit("%s: non-numeric cells %r" % (sheet, odd[:5]))
    return out


# ------------------------------------------------------------------ signs

def norm(name):
    """Wang writes AS1 and CA1 where Cook writes AS01 and CA01."""
    return re.sub(r"(?<=[A-Za-z])0+(?=\d)", "", name.strip())


def parse_transmitters(cells):
    got = set()
    for c in cells:
        t = clean(c).lstrip("*").strip()
        if not t or "uptake" in t:
            continue          # uptake is clearance, not release
        head = re.split(r"[\s(]", t)[0]
        if head in ("ACh", "Glu", "GABA"):
            got.add(head)
    return got


def wang_table(path, first_nt_col, n_nt_cols):
    """{normalised neuron name: (transmitter set, verbatim text)}."""
    rows = rows_of(path, rows_of_first_sheet(path))
    out = {}
    for r in rows[4:]:
        if len(r) <= first_nt_col or not clean(r[2]):
            continue
        cells = [r[first_nt_col + k] for k in range(n_nt_cols) if first_nt_col + k < len(r)]
        text = "; ".join(clean(c) for c in cells if clean(c))
        if not text:
            continue
        names = clean(r[2])
        m = re.fullmatch(r"([A-Z]+)(\d)/(\d)", names)     # DX1/2, EF3/4
        expand = [m.group(1) + m.group(2), m.group(1) + m.group(3)] if m else [names]
        for nm in expand:
            out[norm(nm)] = (parse_transmitters(cells), text)
    return out


def rows_of_first_sheet(path):
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True).sheetnames[0]


# Supplementary file 4 lists sexually dimorphic transmitter use in sex-shared
# neurons. The only entry that changes a classical transmitter is AIM:
# "eat-4 downregulated; unc-17(+)" in the male, i.e. ACh instead of Glu. The
# others add unc-47 alone (not GABA synthesis) or change brightness. Read from
# the file and checked, so a later version of the table cannot slip past.
MALE_OVERRIDES = {"AIM": {"ACh"}}


def male_overrides(path):
    rows = rows_of(path, rows_of_first_sheet(path))
    seen = {clean(r[1]): clean(r[3]) for r in rows if len(r) > 3 and clean(r[1])}
    if "unc-17(+)" not in seen.get("AIM", ""):
        raise SystemExit("Supplementary file 4 no longer says AIM is unc-17(+) in "
                         "the male; re-read it before exporting")
    return MALE_OVERRIDES


# ------------------------------------------------------------------ export

def export(sex, raw, out_dir):
    species = "c_elegans_herm" if sex == "herm" else "c_elegans_male"
    label = "hermaphrodite" if sex == "herm" else "male"
    cells = cell_list(os.path.join(ROOT, raw["cell_lists"]["path"]), sex)
    names = [c[0] for c in cells]
    index = {nm: i for i, nm in enumerate(names)}
    n = len(names)

    # Signs.
    nt = wang_table(os.path.join(ROOT, raw["herm_nt"]["path"]), 20, 3)
    if sex == "male":
        nt_male = wang_table(os.path.join(ROOT, raw["male_nt"]["path"]), 21, 3)
        nt.update(nt_male)
        for cls, tx in male_overrides(os.path.join(ROOT, raw["dimorphic_nt"]["path"])).items():
            for nm in names:
                if re.fullmatch(cls + r"[LR]?", nm):
                    nt[norm(nm)] = (set(tx), "male override from Supplementary file 4")
    signs, no_entry, basis = [], [], {}
    for nm in names:
        e = nt.get(norm(nm))
        if e is None:
            no_entry.append(nm)
            signs.append(Sign.UNKNOWN)
            continue
        s = transmitter_sign(e[0])
        signs.append(s)
        basis[nm] = e[1]
    n_exc = sum(s is Sign.EXCITATORY for s in signs)
    n_inh = sum(s is Sign.INHIBITORY for s in signs)
    unknown_names = [nm for nm, s in zip(names, signs) if s is Sign.UNKNOWN]

    # Chemical synapses.
    adj = os.path.join(ROOT, raw["adjacency"]["path"])
    chem = matrix(adj, "%s chemical" % label)
    unknown_rows = sorted({a for a, _ in chem if a not in index})
    if unknown_rows:
        raise SystemExit("%s chemical: presynaptic rows that are not neurons in SI 4: %s"
                         % (label, unknown_rows))
    synapses, dropped_sign, dropped_sign_w, to_end = [], 0, 0, 0
    nonint = 0
    for (a, b), v in chem.items():
        if b not in index:
            to_end += 1
            continue
        if v != int(v):
            nonint += 1
        s = signs[index[a]]
        if s is Sign.UNKNOWN:
            dropped_sign += 1
            dropped_sign_w += int(v)
            continue
        synapses.append((index[a], index[b], int(round(v)) * s.value))
    if nonint:
        raise SystemExit("%s: %d non-integer chemical weights" % (label, nonint))

    # Gap junctions, from the symmetric table.
    gap = matrix(adj, "%s gap jn symmetric" % label)
    asym = sum(1 for (a, b), v in gap.items() if gap.get((b, a)) != v)
    pairs_all = {tuple(sorted(k)) for k in gap if k[0] != k[1]}
    gaps, self_gap, gap_end = [], 0, 0
    for (a, b) in sorted(pairs_all):
        if a not in index or b not in index:
            gap_end += 1
            continue
        i, j = sorted((index[a], index[b]))
        w = max(gap.get((a, b), 0), gap.get((b, a), 0))
        gaps.append((i, j, float(w)))
    self_gap = sum(1 for (a, b) in gap if a == b)

    # Validation against the paper.
    paper = PAPER[sex]
    chem_edges = len(chem)
    # The paper's count includes the diagonal (junctions between two processes
    # of one cell): 1,433 + 17 = 1,450 in the hermaphrodite sheet.
    gap_edges = len(pairs_all) + self_gap
    report = {
        "neurons": {"got": n, "paper": paper["neurons"]},
        "chemical_edges_all_cells": {"got": chem_edges, "paper": paper["chem_edges"]},
        "gap_edges_all_cells_incl_self": {"got": gap_edges, "paper": paper["gap_edges"]},
        "tolerance": "neurons exact; edge totals within %.0f%% (SI 5 was corrected in "
                     "July 2020, after the paper's counts)" % (100 * EDGE_TOLERANCE),
    }
    fails = []
    if n != paper["neurons"]:
        fails.append("neurons %d != %d" % (n, paper["neurons"]))
    for k, got, want in (("chemical", chem_edges, paper["chem_edges"]),
                         ("gap", gap_edges, paper["gap_edges"])):
        if abs(got - want) > EDGE_TOLERANCE * want:
            fails.append("%s edges %d vs paper %d" % (k, got, want))

    # Sensory.
    sensory, senses, sens_prov = {}, {}, {}
    notes = {c[0]: (c[1], c[2]) for c in cells}
    for ch, spec in SENSORY.items():
        for nm in spec["neurons"]:
            typ, note = notes.get(nm, ("", ""))
            if typ != spec["annotation"][0] or spec["annotation"][1] not in note:
                fails.append("%s: SI 4 annotates %s as %r / %r, not %r"
                             % (ch, nm, typ, note, spec["annotation"]))
        sensory[ch] = [index[nm] for nm in spec["neurons"] if nm in index]
        senses[ch] = spec["signal"]
        sens_prov[ch] = {"signal": spec["signal"], "neurons": spec["neurons"],
                         "why": spec["why"]}

    # Motor.
    motor_names = [c[0] for c in cells if c[1] == "motorneuron"]
    units = [[index[x] for x in u] for u in pk.bilateral_units(motor_names)]
    groups, rule = pk.deal_groups(units)
    motor = sorted({i for m in pk.MODULES for g in groups[m] for i in g})

    meta = {
        "species": species,
        "dataset": "Cook et al. (2019)",
        "version": "SI 5 Connectome adjacency matrices, corrected July 2020 "
                   "(WormWiring.org); SI 4 Cell lists",
        "n": n,
        "names": names,
        "params": dict(pk.SHIU_PARAMS),
        "sensory": sensory,
        "senses": senses,
        "groups": groups,
        "motor": motor,
        "partition_seed": pk.PARTITION_SEED,
        "window_ms": 200,
        "provenance": {
            "neurons": ("Cook et al. 2019 SI 4 'Cell lists': every cell typed sensory, "
                        "interneuron, motorneuron or neuron on the pharynx, sex-shared and "
                        "%s sheets, in sheet order (%d). Muscles, glia and other end "
                        "organs are not neurons and are not nodes." % (
                            "hermaphrodite specific" if sex == "herm" else "male-specific", n)),
            "chemical": ("SI 5 '%s chemical', rows presynaptic. Weight = Cook's number of "
                         "EM serial sections of connectivity (synapse number and size, with "
                         "extrapolated connections), stored as the i16 count, so w_syn "
                         "multiplies sections, not synapses. %d neuron->neuron edges kept; "
                         "%d edges onto muscles/end organs left out." % (label, len(synapses), to_end)),
            "gap": ("SI 5 '%s gap jn symmetric', neuron pairs stored once (a < b), "
                    "weight in the same serial-section unit. %d kept; %d pairs with a "
                    "muscle or end organ left out; %d self-junction cells skipped; %d "
                    "cells of the symmetric sheet differ from their mirror (the larger "
                    "is kept)." % (label, len(gaps), gap_end, self_gap, asym)),
            "signs": ("Wang et al. 2024 (eLife 13:RP95402) Supplementary file 2%s, column "
                      "'Neurotransmitter(s)' (all three cells of it), through "
                      "open_animal_stage.signs.transmitter_sign: ACh or Glu excitatory, "
                      "GABA inhibitory, both kinds or neither UNKNOWN. '(uptake)' entries "
                      "are clearance, not release, and are ignored; dim-and-variable (*) "
                      "entries count. GLUTAMATE IS EXCITATORY EVERYWHERE by rule, although "
                      "GluCl channels make some glutamatergic synapses inhibitory (e.g. AWC "
                      "-> AIY); per-synapse receptor data is not used. Result: %d "
                      "excitatory, %d inhibitory, %d unknown; %d neurons had no row in the "
                      "atlas: %s." % (
                          " (hermaphrodite) and 3 (male-specific neurons), with the AIM male "
                          "override from Supplementary file 4" if sex == "male" else "",
                          n_exc, n_inh, len(unknown_names), len(no_entry),
                          ", ".join(no_entry) or "none")),
            "unknown_sign_neurons": unknown_names,
            "sensory": sens_prov,
            "motor": ("SI 4 cell type 'motorneuron' (pharyngeal, head, sublateral, ventral "
                      "cord%s): %d neurons, as bilateral units (XXXL+XXXR together, "
                      "everything else alone; %d units), dealt by Open Fly's make_groups "
                      "(numpy default_rng(783), 'balanced' rule) into 12+12+8+7 groups. "
                      "The pharyngeal motor neurons are included because the source "
                      "annotates them motor; the pharynx connects to the rest only "
                      "through RIP, so those groups will be quiet." % (
                          ", HSN and VC" if sex == "herm" else "; SI 4 annotates no "
                          "male-specific neuron as motor, so CA/CP are not here",
                          len(motor_names), len(units))),
            "motor_rule": rule,
            "params": ("Open Fly's constants from Shiu et al. 2024 (fly LIF), unchanged; "
                       "w_gap 0.0 and calibrated false -- w_syn and w_gap are set later by "
                       "one fixed rule in the JS brain, not tuned here. Most C. elegans "
                       "neurons are graded (non-spiking), so a leaky integrate-and-fire "
                       "neuron is a stated simplification, not a model of the worm's "
                       "physiology."),
            "licence": LICENCE,
            "dropped_unknown_sign": dropped_sign,
            "dropped_unknown_sign_weight": dropped_sign_w,
            "validation": report,
            "raw": raw,
            "exported": datetime.date.today().isoformat(),
            "exporter": "tools/export_celegans.py",
        },
    }
    if fails:
        raise SystemExit("%s: validation failed:\n  %s" % (species, "\n  ".join(fails)))
    pk.write_package(out_dir, n, synapses, gaps, meta)
    bad = pk.check_package(out_dir)
    if bad:
        raise SystemExit("%s: package check failed:\n  %s" % (species, "\n  ".join(bad)))

    print("%s: %d neurons (paper %d), %d chemical kept of %d neuron->neuron "
          "(%d dropped for unknown sign), %d gap junctions; all-cell edges "
          "chem %d vs paper %d, gap %d vs paper %d; signs +%d -%d ?%d; "
          "motor %d in %s groups"
          % (species, n, paper["neurons"], len(synapses), len(synapses) + dropped_sign,
             dropped_sign, len(gaps), chem_edges, paper["chem_edges"], gap_edges,
             paper["gap_edges"], n_exc, n_inh, len(unknown_names), len(motor), rule))
    return {
        "dataset_version": meta["version"],
        "exported": meta["provenance"]["exported"],
        "exporter": "tools/export_celegans.py",
        "package": os.path.relpath(out_dir, ROOT),
        "sources": raw,
        "licence": LICENCE,
        "counts": {"neurons": n, "chemical_synapses": len(synapses),
                   "gap_junctions": len(gaps), "dropped_unknown_sign": dropped_sign,
                   "motor_neurons": len(motor)},
        "validation": report,
    }


def update_roster(blocks):
    path = os.path.join(ROOT, "data", "roster.json")
    d = json.load(open(path))
    for s in d["species"]:
        if s["id"] in blocks:
            s["export"] = blocks[s["id"]]
    with open(path + ".tmp", "w") as f:
        json.dump(d, f, indent=2, ensure_ascii=False)
        f.write("\n")
    os.replace(path + ".tmp", path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-fetch", action="store_true")
    ap.add_argument("--no-roster", action="store_true",
                    help="do not rewrite the export blocks in data/roster.json")
    args = ap.parse_args()
    raw = raw_files(not args.no_fetch)
    blocks = {}
    for sex, sid in (("herm", "c_elegans_herm"), ("male", "c_elegans_male")):
        blocks[sid] = export(sex, raw, os.path.join(ROOT, "web", "species", sid))
    if not args.no_roster:
        update_roster(blocks)


if __name__ == "__main__":
    main()
